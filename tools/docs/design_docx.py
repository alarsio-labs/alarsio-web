# -*- coding: utf-8 -*-
"""Renders a content module to an editable DOCX in docs/.

    python tools/docs/design_docx.py                   -> the technical design document
    python tools/docs/design_docx.py roadmap_content   -> the product feature roadmap

A4 landscape, real Word tables, and each diagram placed as a 2x PNG.
design_pdf.py renders the same content module to PDF, so a document's two
formats never drift.

Needs python-docx. Diagrams are rasterised with the same headless Chrome the
other generators use.
"""
import io
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib   # noqa: E402

# Which document to build. Defaults to the technical design document;
# pass another content module to build a different one.
C = importlib.import_module(sys.argv[1] if len(sys.argv) > 1 else "design_content")
from paths import DOCS, BUILD, chrome   # noqa: E402

OUT = os.path.join(DOCS, C.OUT_STEM + ".docx")

FOREST = RGBColor(0x12, 0x37, 0x2A)
GOLD = RGBColor(0xC9, 0xA6, 0x48)
INK = RGBColor(0x16, 0x26, 0x1F)
LEAD = RGBColor(0x3C, 0x51, 0x47)
MUTED = RGBColor(0x5F, 0x71, 0x68)
CODE = RGBColor(0x33, 0x51, 0x3F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x8A, 0x97, 0x90)

BODY_FONT = "Segoe UI"
HEAD_FONT = "Georgia"
MONO_FONT = "Consolas"

TEXT_MM = 269.0          # 297 page - 14 - 14 margins
CONTENT = Mm(TEXT_MM)


# ── low-level OOXML helpers ──────────────────────────────────────────────────
# Word validates the order of the children of pPr, rPr, tcPr and tblPr. An
# element appended at the end is rejected as a corrupt file, so every raw
# element below is inserted at its schema position.
ORDER = {
    "pPr": ("pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
            "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap "
            "overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
            "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
            "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr "
            "sectPr pPrChange").split(),
    "rPr": ("rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
            "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz "
            "szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang "
            "eastAsianLayout specVanish oMath").split(),
    "tcPr": ("cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
             "tcFitText vAlign hideMark headers cellIns cellDel cellMerge tcPrChange").split(),
    "tblPr": ("tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
              "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook "
              "tblCaption tblDescription tblPrChange").split(),
}


def _name(el):
    return el.tag.split("}")[-1]


def put(parent, el):
    """Insert el among parent's children at its schema position, replacing any
    existing element with the same name."""
    order = ORDER[_name(parent)]
    name = _name(el)
    for child in list(parent):
        if _name(child) == name:
            parent.remove(child)
    idx = order.index(name)
    for child in parent:
        cn = _name(child)
        if cn in order and order.index(cn) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def shade(obj, fill):
    """Background fill on a cell or a run."""
    pr = obj._tc.get_or_add_tcPr() if hasattr(obj, "_tc") else obj._r.get_or_add_rPr()
    put(pr, _el("w:shd", val="clear", color="auto", fill=fill))


def borders(el_pr, spec, tag):
    """spec: {'left': (size_eighths, colour)} ; sides left out get 'nil'."""
    b = OxmlElement(tag)
    for side in ("top", "left", "bottom", "right"):
        s = OxmlElement("w:" + side)
        if side in spec:
            size, colour = spec[side]
            s.set(qn("w:val"), "single")
            s.set(qn("w:sz"), str(size))
            s.set(qn("w:color"), colour)
        else:
            s.set(qn("w:val"), "nil")
        s.set(qn("w:space"), "0")
        b.append(s)
    put(el_pr, b)


def para_border(p, spec):
    borders(p._p.get_or_add_pPr(), spec, "w:pBdr")


def cell_border(cell, spec):
    borders(cell._tc.get_or_add_tcPr(), spec, "w:tcBorders")


def spacing(run, twentieths):
    """Letter spacing, in twentieths of a point."""
    put(run._r.get_or_add_rPr(), _el("w:spacing", val=str(twentieths)))


def cell_margins(tbl, top, left, bottom, right):
    m = OxmlElement("w:tblCellMar")
    for side, mm in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(int(mm * 56.7)))   # mm -> twips
        e.set(qn("w:type"), "dxa")
        m.append(e)
    put(tbl._tbl.tblPr, m)


def keep_together(p):
    put(p._p.get_or_add_pPr(), _el("w:keepNext", val="true"))


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def field(p, instr):
    """A Word field, e.g. PAGE."""
    r = p.add_run()
    r._r.append(_el("w:fldChar", fldCharType="begin"))
    r2 = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    r3._r.append(_el("w:fldChar", fldCharType="end"))
    for r_ in (r, r2, r3):
        r_.font.name = BODY_FONT
        r_.font.size = Pt(7.5)
        r_.font.color.rgb = RGBColor(0x9A, 0xA4, 0x9E)


# ── inline markup ────────────────────────────────────────────────────────────
TOKEN = re.compile(r"(`[^`]+`|~[^~]+~)")


def runs(p, text, size, colour=INK, bold=False, font=BODY_FONT):
    """Write text into paragraph p, honouring `code` and ~tag~ markup."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(size * 0.92)
            r.font.color.rgb = CODE
            shade(r, "F3F1E9")
        elif part.startswith("~") and part.endswith("~"):
            r = p.add_run(" " + part[1:-1])
            r.font.name = BODY_FONT
            r.font.size = Pt(size * 0.86)
            r.font.color.rgb = GREY
        else:
            r = p.add_run(part)
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = colour
            r.font.bold = bold
    return p


# ── diagram rasterising ──────────────────────────────────────────────────────
CHROME = chrome()


def raster(svg, name):
    """SVG string -> 2x PNG. Returns (path, width_px, height_px)."""
    vb = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg)
    w, h = int(float(vb.group(1))), int(float(vb.group(2)))
    svg_path = os.path.join(BUILD, name + ".svg")
    io.open(svg_path, "w", encoding="utf-8").write(svg)
    html = os.path.join(BUILD, "_r.html")
    io.open(html, "w", encoding="utf-8").write(
        '<!doctype html><meta charset="utf-8">'
        '<style>html,body{margin:0;padding:0;background:#fff}'
        'img{display:block;width:%dpx;height:%dpx}</style><img src="file:///%s">'
        % (w, h, svg_path.replace("\\", "/")))
    png = os.path.join(BUILD, name + ".png")
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-sandbox",
                    "--hide-scrollbars", "--force-device-scale-factor=2",
                    "--window-size=%d,%d" % (w, h), "--screenshot=" + png,
                    "file:///" + html.replace("\\", "/")], capture_output=True)
    if not os.path.exists(png):
        raise SystemExit("Chrome did not write " + png)
    return png, w, h


# ── block renderers ──────────────────────────────────────────────────────────
def table_width(tbl, width_mm):
    put(tbl._tbl.tblPr, _el("w:tblW", w=str(int(width_mm * 56.7)), type="dxa"))


def add_table(host, cols, rows, widths, size, width_mm=TEXT_MM):
    fs = 7.0 if size == "xs" else 7.5
    tbl = host.add_table(rows=len(rows) + 1, cols=len(cols))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    put(tbl._tbl.tblPr, _el("w:tblLayout", type="fixed"))
    cell_margins(tbl, 1.1, 2.0, 1.1, 2.0)
    table_width(tbl, width_mm)
    widths_mm = [width_mm * float(w.rstrip("%")) / 100.0 for w in widths]

    for j, (c, wmm) in enumerate(zip(cols, widths_mm)):
        cell = tbl.cell(0, j)
        cell.width = Mm(wmm)
        shade(cell, "12372A")
        cell_border(cell, {s: (6, "12372A") for s in ("top", "left", "bottom", "right")})
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        runs(p, c, fs + 0.4, WHITE, bold=True)

    for i, row in enumerate(rows, start=1):
        for j, (val, wmm) in enumerate(zip(row, widths_mm)):
            cell = tbl.cell(i, j)
            cell.width = Mm(wmm)
            if i % 2 == 0:
                shade(cell, "FAF9F4")
            cell_border(cell, {s: (4, "DED8CA") for s in ("top", "left", "bottom", "right")})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            runs(p, val, fs)
    return tbl


def panel_table(host, ncols, fills=("F6F8F5",), accent="4F9A72"):
    tbl = host.add_table(rows=1, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    put(tbl._tbl.tblPr, _el("w:tblLayout", type="fixed"))
    cell_margins(tbl, 1.6, 2.4, 1.6, 2.4)
    return tbl


def add_rulebar(host, items, ncols, width_mm=TEXT_MM):
    tbl = panel_table(host, ncols)
    table_width(tbl, width_mm)
    w = Mm(width_mm / ncols)
    for j, (head, rest) in enumerate(items):
        cell = tbl.cell(0, j)
        cell.width = w
        shade(cell, "F6F8F5")
        cell_border(cell, {"left": (18, "4F9A72"), "top": (4, "DFE5DE"),
                           "bottom": (4, "DFE5DE"), "right": (4, "DFE5DE")})
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.16
        runs(p, head + " ", 8, FOREST, bold=True)
        runs(p, rest, 8)
    return tbl


def add_note(host, text, width_mm=TEXT_MM):
    tbl = panel_table(host, 1)
    table_width(tbl, width_mm)
    cell = tbl.cell(0, 0)
    cell.width = Mm(width_mm)
    shade(cell, "F6F4EC")
    cell_border(cell, {"left": (24, "C9A648"), "top": (4, "E2DCCA"),
                       "bottom": (4, "E2DCCA"), "right": (4, "E2DCCA")})
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    runs(p, text, 8)
    return tbl


def add_phases(host, phases, width_mm):
    tbl = host.add_table(rows=len(phases), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    put(tbl._tbl.tblPr, _el("w:tblLayout", type="fixed"))
    cell_margins(tbl, 1.4, 2.4, 1.4, 2.4)
    table_width(tbl, width_mm)
    for i, (num, title, mods) in enumerate(phases):
        left, right = tbl.cell(i, 0), tbl.cell(i, 1)
        left.width, right.width = Mm(width_mm * 0.20), Mm(width_mm * 0.80)
        for cell in (left, right):
            shade(cell, "F6F8F5")
        cell_border(left, {"left": (18, "4F9A72"), "top": (4, "DFE5DE"),
                           "bottom": (4, "DFE5DE")})
        cell_border(right, {"top": (4, "DFE5DE"), "bottom": (4, "DFE5DE"),
                            "right": (4, "DFE5DE")})
        p = left.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(num.upper())
        r.font.name = BODY_FONT
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = GOLD
        spacing(r, 24)

        p = right.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        runs(p, title, 8.5, FOREST, bold=True)
        p2 = right.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(0)
        runs(p2, mods, 7.2, MUTED, font=MONO_FONT)
    return tbl


def add_rules(host, rules):
    for i, rule in enumerate(rules, start=1):
        p = host.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(3.4)
        pf.left_indent = Mm(6)
        pf.first_line_indent = Mm(-6)
        pf.line_spacing = 1.16
        runs(p, "%d.  " % i, 8, MUTED)
        runs(p, rule, 8)


def add_h3(host, text):
    p = host.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    keep_together(p)
    runs(p, text, 11.5, FOREST, bold=True, font=HEAD_FONT)


def add_diagram(host, svg, name, scale, width_mm=TEXT_MM):
    png, w, h = raster(svg, name)
    p = host.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(png, width=Mm(width_mm * scale))


def gap(host, pts=5):
    p = host.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pts)
    return p


def render(host, blocks, idx, width_mm=TEXT_MM):
    """Render a list of content blocks into a document or a table cell."""
    for n, b in enumerate(blocks):
        kind = b[0]
        if kind == "diagram":
            scale = b[3] if len(b) > 3 else 1.0
            add_diagram(host, b[1], "d_%s_%d" % (idx, n), scale, width_mm)
        elif kind == "table":
            add_table(host, b[1], b[2], b[3], b[4], width_mm)
            gap(host)
        elif kind == "rulebar":
            add_rulebar(host, b[1], b[2], width_mm)
            gap(host)
        elif kind == "note":
            add_note(host, b[1], width_mm)
            gap(host)
        elif kind == "h3":
            add_h3(host, b[1])
        elif kind == "rules":
            add_rules(host, b[1])
        elif kind == "phases":
            add_phases(host, b[1], width_mm)
            gap(host)
        elif kind == "twocol":
            outer = host.add_table(rows=1, cols=2)
            outer.alignment = WD_TABLE_ALIGNMENT.LEFT
            outer.autofit = False
            put(outer._tbl.tblPr, _el("w:tblLayout", type="fixed"))
            cell_margins(outer, 0, 0, 0, 3.5)
            table_width(outer, width_mm)
            half = width_mm / 2.0
            for cell, sub in ((outer.cell(0, 0), b[1]), (outer.cell(0, 1), b[2])):
                cell.width = Mm(half)
                cell_border(cell, {})
                render(cell, sub, idx, half - 3.5)
                tidy_cell(cell)
        else:
            raise ValueError(kind)


def tidy_cell(cell):
    """Drop the placeholder paragraph a cell opens with. Never touch the last
    child: Word requires every cell to end in a paragraph."""
    kids = list(cell._tc)
    body = [k for k in kids if _name(k) in ("p", "tbl")]
    if len(body) > 1 and _name(body[0]) == "p" and not "".join(body[0].itertext()).strip():
        cell._tc.remove(body[0])


# ── document ─────────────────────────────────────────────────────────────────
doc = Document()

style = doc.styles["Normal"]
style.font.name = BODY_FONT
style.font.size = Pt(9)
style.font.color.rgb = INK
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = 1.2
style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Mm(297), Mm(210)
sec.left_margin = sec.right_margin = Mm(14)
sec.top_margin, sec.bottom_margin = Mm(13), Mm(14)
sec.footer_distance = Mm(7)

# The built-in Footer style carries a centre tab, which would swallow the
# first tab and park the page number mid-page. Clear the style's stops.
doc.styles["Footer"].paragraph_format.tab_stops.clear_all()
fp = sec.footer.paragraphs[0]
fp.paragraph_format.tab_stops.clear_all()
fp.paragraph_format.tab_stops.add_tab_stop(CONTENT, WD_TAB_ALIGNMENT.RIGHT)
para_border(fp, {"top": (4, "E6E1D3")})
fr = fp.add_run(C.FOOTER + "\t")
fr.font.name = BODY_FONT
fr.font.size = Pt(7.5)
fr.font.color.rgb = RGBColor(0x9A, 0xA4, 0x9E)
field(fp, "PAGE")

# ── cover ────────────────────────────────────────────────────────────────────
cv = C.COVER

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run(cv["eyebrow"].upper())
r.font.name = BODY_FONT
r.font.size = Pt(9)
r.font.bold = True
r.font.color.rgb = RGBColor(0x0F, 0x7A, 0x45)
spacing(r, 48)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run(cv["title"])
r.font.name = HEAD_FONT
r.font.size = Pt(34)
r.font.bold = True
r.font.color.rgb = FOREST

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run(cv["title_em"])
r.font.name = HEAD_FONT
r.font.size = Pt(25)
r.font.bold = True
r.font.color.rgb = GOLD

rule = doc.add_table(rows=1, cols=1)
rule.autofit = False
put(rule._tbl.tblPr, _el("w:tblLayout", type="fixed"))
rc = rule.cell(0, 0)
rc.width = Mm(48)
cell_border(rc, {"top": (30, "C9A648")})
rc.paragraphs[0].paragraph_format.space_after = Pt(0)
rc.paragraphs[0].add_run("").font.size = Pt(2)
gap(doc, 7)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.paragraph_format.line_spacing = 1.4
runs(p, cv["lede"], 11, LEAD, font=HEAD_FONT)

stats = doc.add_table(rows=1, cols=len(cv["stats"]))
stats.autofit = False
put(stats._tbl.tblPr, _el("w:tblLayout", type="fixed"))
cell_margins(stats, 2.0, 0, 0, 5.0)
for j, (big, label, sub) in enumerate(cv["stats"]):
    cell = stats.cell(0, j)
    cell.width = Mm(TEXT_MM / len(cv["stats"]))
    cell_border(cell, {"top": (12, "CFC7B5")})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(big)
    r.font.name = HEAD_FONT
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = FOREST
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    runs(p2, label + "\n" + sub, 8, RGBColor(0x6B, 0x7A, 0x72))
gap(doc, 6)

add_h3(doc, "Contents")
toc = doc.add_table(rows=len(cv["toc"]), cols=2)
toc.autofit = False
put(toc._tbl.tblPr, _el("w:tblLayout", type="fixed"))
cell_margins(toc, 0.8, 0, 0.8, 2.0)
for i, (title, num) in enumerate(cv["toc"]):
    a, b_ = toc.cell(i, 0), toc.cell(i, 1)
    a.width, b_.width = Mm(TEXT_MM - 20), Mm(20)
    cell_border(a, {"bottom": (4, "DFE5DE")})
    cell_border(b_, {"bottom": (4, "DFE5DE")})
    pa = a.paragraphs[0]
    pa.paragraph_format.space_after = Pt(0)
    runs(pa, "%s.  " % (i + 1), 8.5, GREY)
    runs(pa, title, 8.5)
    pb = b_.paragraphs[0]
    pb.paragraph_format.space_after = Pt(0)
    pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    runs(pb, num, 8.5, GOLD, bold=True)

# ── body pages ───────────────────────────────────────────────────────────────
for pg in C.PAGES:
    page_break(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    keep_together(p)
    r = p.add_run(pg["kicker"].upper())
    r.font.name = BODY_FONT
    r.font.size = Pt(7.5)
    r.font.bold = True
    r.font.color.rgb = GOLD
    spacing(r, 44)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    para_border(p, {"bottom": (16, "12372A")})
    keep_together(p)
    r = p.add_run(pg["title"])
    r.font.name = HEAD_FONT
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = FOREST

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.34
    keep_together(p)
    runs(p, pg["lead"], 10.5, LEAD, font=HEAD_FONT)

    render(doc, pg["blocks"], pg["num"])

doc.save(OUT)
print("docx:", os.path.exists(OUT), os.path.getsize(OUT))
